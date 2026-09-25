#!/usr/bin/env python3
"""Extrai texto de arquivos de projeto (docx, pptx, xlsx, pdf, txt, md, eml) para leitura na etapa de descoberta.

Uso:
    python3 scripts/extrair_texto.py <arquivo> [<arquivo> ...]
    python3 scripts/extrair_texto.py projetos/<slug>/entrada/      (todos os arquivos da pasta)

Dependências opcionais: python-docx, python-pptx, openpyxl, pdfplumber (ou pdftotext no sistema).
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path


def docx_text(p: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = Document(p)
    out = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, d).text.strip()
            if t:
                out.append(t)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                cells = []
                for c in row.cells:
                    txt = c.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != txt:
                        cells.append(txt)
                out.append("| " + " | ".join(cells) + " |")
            out.append("")
    return "\n".join(out)


def pptx_text(p: Path) -> str:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(p).slides, 1):
        out.append(f"--- slide {i} ---")
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text.strip())
    return "\n".join(out)


def xlsx_text(p: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- aba {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append("| " + " | ".join("" if v is None else str(v) for v in row) + " |")
    return "\n".join(out)


def pdf_text(p: Path) -> str:
    try:
        import pdfplumber

        with pdfplumber.open(p) as pdf:
            return "\n".join((pg.extract_text() or "") for pg in pdf.pages)
    except ImportError:
        return subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True).stdout


def extract(p: Path) -> str:
    ext = p.suffix.lower()
    try:
        if ext == ".docx":
            return docx_text(p)
        if ext == ".pptx":
            return pptx_text(p)
        if ext in (".xlsx", ".xlsm"):
            return xlsx_text(p)
        if ext == ".pdf":
            return pdf_text(p)
        return p.read_text(encoding="utf-8", errors="replace")
    except ImportError as e:
        return f"[não foi possível ler {p.name}: dependência ausente ({e.name}). Instale e rode de novo.]"


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    paths: list[Path] = []
    for a in sys.argv[1:]:
        p = Path(a)
        paths.extend(sorted(x for x in p.iterdir() if x.is_file()) if p.is_dir() else [p])
    for p in paths:
        print(f"\n===== {p} =====\n")
        print(extract(p))
    return 0


if __name__ == "__main__":
    sys.exit(main())
