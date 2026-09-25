#!/usr/bin/env python3
"""Gera o .docx no padrão visual okno a partir do formato intermediário (.md).

Uso:
    python3 scripts/gerar_docx.py projetos/<slug>/<arquivo>.md [saida.docx]

Requer: pip install python-docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("python-docx não instalado. Rode: pip install python-docx")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from okno_md import QID_RE, SECTION_RE, is_sep_row, parse_frontmatter, split_row  # noqa: E402

TEAL = "00808A"
GRAY = "D9D9D9"
LIGHT = "F2F2F2"
BLUE = RGBColor(0x1F, 0x4E, 0x9A)     # programação
RED = RGBColor(0xC0, 0x00, 0x00)      # entrevistador / moderador
DARK = RGBColor(0x40, 0x40, 0x40)
FONT = "Arial"
FOOTER = "Rua Aspicuelta, 422 - conjunto 34A – Vila Madalena - São Paulo | SP   www.okno-ne.com.br"

PROG_PREFIX = ("PROGRAMAÇÃO", "PROGRAMACAO", "PROGRAMADOR", "PROGRAMMING", "PROG:", "PN:", "PN ", "PROCESSAMENTO",
               "APLICAR", "APPLY", "SOMENTE", "ONLY", "DEVE ", "MUST ", "SHOW ", "MOSTRAR ", "PERGUNTAR ",
               "REPETIR", "REPEAT", "ATENÇÃO", "SE ", "IF ")
FIELD_PREFIX = ("ENTREVISTADOR", "INTERVIEWER", "MODERADOR", "MODERATOR")


# ---------------------------------------------------------------- helpers de formatação
def shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    old = tcPr.find(qn("w:shd"))
    if old is not None:
        tcPr.remove(old)
    # ordem do schema: shd antes de noWrap/tcMar/textDirection/tcFitText/vAlign/hideMark
    anchor = None
    for tag in ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"):
        anchor = tcPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(shd)
    else:
        tcPr.append(shd)


def set_borders(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    # ordem do schema OOXML: tblBorders vem antes de shd/tblLayout/tblCellMar/tblLook
    anchor = None
    for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"):
        anchor = tblPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        tblPr.append(borders)


def add_runs(par, text: str, size: float = 10, bold: bool = False, color: RGBColor | None = None,
             italic: bool = False) -> None:
    """Suporta **negrito** inline."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        b = bold
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        run = par.add_run(part)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = b
        run.italic = italic
        if color is not None:
            run.font.color.rgb = color


def para(doc, text: str = "", size: float = 10, bold: bool = False, color=None, align=None,
         space_after: float = 4, italic: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    if text:
        add_runs(p, text, size=size, bold=bold, color=color, italic=italic)
    return p


def teal_box(doc, text: str, size: float = 11) -> None:
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, TEAL)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(c.paragraphs[0], text.upper(), size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    para(doc, "", space_after=2)


def cell_text(cell, text: str, size: float = 9, bold: bool = False, center: bool = False) -> None:
    gray = False
    if text.startswith("[CINZA]"):
        text, gray = text[len("[CINZA]"):].strip(), True
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    up = text.upper()
    color = None
    if up in ("ENCERRE", "TERMINATE") or up.startswith(("PULE", "SKIP", "VÁ PARA", "GO TO")):
        color, bold = RED, True
    elif up in ("CONTINUE",):
        bold = True
    add_runs(p, text, size=size, bold=bold, color=color)
    if gray:
        shade(cell, GRAY)


def md_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], [r for r in rows[1:] if not is_sep_row(r)]
    ncols = max(len(header), *(len(r) for r in body)) if body else len(header)
    hide_header = all(h.strip().lower() in ("", "opção", "opcao", "option", "cód", "cod", "código", "ação", "acao",
                                            "action") for h in header)
    t = doc.add_table(rows=0, cols=ncols)
    set_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if not hide_header:
        cells = t.add_row().cells
        for i in range(ncols):
            txt = header[i] if i < len(header) else ""
            cell_text(cells[i], txt, bold=True, center=i > 0)
            shade(cells[i], LIGHT)
    for r in body:
        cells = t.add_row().cells
        for i in range(ncols):
            txt = r[i] if i < len(r) else ""
            cell_text(cells[i], txt, center=i > 0)
    para(doc, "", space_after=4)


# ---------------------------------------------------------------- blocos fixos do documento
def header_block(doc, meta: dict) -> None:
    tipo = str(meta.get("tipo", ""))
    en = str(meta.get("idioma", "pt")).lower().startswith("en")
    t = doc.add_table(rows=2, cols=5)
    set_borders(t)
    r0 = ["okno", meta.get("painel", ""), "START __:__" if en else "INÍCIO __:__",
          "END __:__" if en else "TÉRMINO __:__",
          ("QUESTIONNAIRE NUMBER" if en else "NÚMERO QUESTIONÁRIO") + " |___|___|___|___|"]
    r1 = [f"JOB {meta.get('job', 'A DEFINIR')}", str(meta.get("estudo", "")).upper(),
          f"{meta.get('entrevistas', '')} {'INTERVIEWS' if en else 'ENTREVISTAS'}".strip()
          if str(meta.get("entrevistas", "")).isdigit() else str(meta.get("entrevistas", "")).upper(),
          str(meta.get("data", "")).upper(), ""]
    for i, v in enumerate(r0):
        cell_text(t.cell(0, i), str(v), bold=i == 0, center=True)
    for i, v in enumerate(r1):
        cell_text(t.cell(1, i), str(v), bold=True, center=True)
    shade(t.cell(0, 0), TEAL)
    t.cell(0, 0).paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    para(doc, "", space_after=2)

    if tipo in ("quali-roteiro", "quali-pre-tarefa", "diario-de-uso"):
        return
    fields = (["RESPONDENT'S NAME:", "HOME PHONE: (   )", "CELLPHONE: (   )", "ADDRESS:", "POSTAL CODE: |__|__|__|__|__|-|__|__|__|",
               "NEIGHBORHOOD:", "E-MAIL:"] if en else
              ["NOME DO ENTREVISTADO:", "FONE RES: (   )      FONE COM: (   )      FONE CEL: (   )", "ENDEREÇO:",
               "CEP: |__|__|__|__|__|-|__|__|__|      BAIRRO:      ZONA:", "E-MAIL:"])
    ft = doc.add_table(rows=len(fields), cols=1)
    set_borders(ft)
    for i, f in enumerate(fields):
        cell_text(ft.cell(i, 0), f, size=8)
    para(doc, "", space_after=2)

    cpf = "|__|__|__|.|__|__|__|.|__|__|__|-|__|__|"
    roles = (["INTERVIEWER:", "CRITIC:", "CHECKER:"] if en else ["ENTREVISTADOR(A):", "CRÍTICO(A):", "VERIFICADOR(A):"])
    ct = doc.add_table(rows=4, cols=3)
    set_borders(ct)
    for i, role in enumerate(roles):
        cell_text(ct.cell(i, 0), role, size=8, bold=True)
        cell_text(ct.cell(i, 1), f"CPF: {cpf}", size=8)
        cell_text(ct.cell(i, 2), "DATE: __/__/____" if en else "DATA: __/__/____", size=8)
    cell_text(ct.cell(3, 0), "ROUNDS OF PROCESSING AND CODIFICATION" if en else "VOLTAS DE PROCESSAMENTO E CODIFICAÇÃO",
              size=8, bold=True)
    cell_text(ct.cell(3, 1), ("ENCODER CODE" if en else "CÓDIGO DO CODIFICADOR") + ": |____|____|____|", size=8)
    para(doc, "", space_after=4)


def footer_block(doc, meta: dict) -> None:
    for section in doc.sections:
        f = section.footer
        p = f.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, FOOTER, size=9, color=DARK)
        if meta.get("classificacao"):
            p2 = f.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p2, str(meta["classificacao"]), size=8, bold=True, color=DARK)


# ---------------------------------------------------------------- conversão
def convert(src: Path, dst: Path) -> None:
    lines = src.read_text(encoding="utf-8").splitlines()
    meta, start = parse_frontmatter(lines)

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(1.8)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10)

    header_block(doc, meta)
    if meta.get("titulo"):
        teal_box(doc, str(meta["titulo"]), size=12)

    tbl: list[list[str]] = []
    in_mapa = False
    after_encerramento = False

    def flush():
        nonlocal tbl
        if tbl:
            md_table(doc, tbl)
            tbl = []

    for raw in lines[start:]:
        s = raw.rstrip()
        if in_mapa:
            break
        if s.strip().startswith("|"):
            tbl.append(split_row(s))
            continue
        flush()
        if not s.strip():
            continue

        m = SECTION_RE.match(s)
        if m and not s.startswith("###"):
            name = m.group(1).strip()
            up = name.upper()
            if up.startswith("MAPA DE COBERTURA"):
                in_mapa = True
                continue
            if up.startswith(("SEÇÃO:", "SECAO:", "SECTION:")):
                teal_box(doc, name.split(":", 1)[1].strip())
            elif up.startswith("ENCERRAMENTO"):
                after_encerramento = True
                para(doc, "TEXTO DE ENCERRAMENTO", bold=True, space_after=2)
            elif up.startswith("INTRODU"):
                para(doc, "", space_after=2)
            elif up.startswith("COTAS"):
                para(doc, "COTAS", bold=True, space_after=2)
            else:
                para(doc, name.upper(), bold=True, space_after=2)
            continue

        mq = QID_RE.match(s)
        if mq:
            label, qtext = mq.group(1).rstrip("."), mq.group(2).strip()
            p = para(doc, "", space_after=3)
            p.paragraph_format.space_before = Pt(4)
            add_runs(p, f"{label}. ", bold=True, size=10)
            add_runs(p, qtext, bold=True, size=10)
            continue
        if s.startswith("###"):
            p = para(doc, s.lstrip("#").strip(), bold=True, size=10.5, space_after=3)
            p.paragraph_format.space_before = Pt(6)
            continue

        st = s.strip()
        up = st.upper()
        if up.startswith(("NOVA TELA", "NEW SCREEN")):
            para(doc, st, size=8, bold=True, color=DARK, space_after=1)
        elif up.startswith(FIELD_PREFIX):
            para(doc, st, size=9, bold=True, color=RED, space_after=2)
        elif up.startswith(PROG_PREFIX) or (up == st and len(st) > 12 and any(c.isalpha() for c in st)):
            para(doc, st, size=9, bold=True, color=BLUE, space_after=2)
        elif st.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, st[2:], size=10)
        else:
            para(doc, st, size=10)
    flush()

    if after_encerramento:
        t = doc.add_table(rows=2, cols=2)
        set_borders(t)
        cell_text(t.cell(0, 0), "VOLTAS DE CAMPO", bold=True, center=True)
        cell_text(t.cell(0, 1), "COMENTÁRIOS DE CAMPO", bold=True, center=True)
    footer_block(doc, meta)
    doc.save(dst)


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    convert(src, dst)
    print(f"OK: {dst}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
